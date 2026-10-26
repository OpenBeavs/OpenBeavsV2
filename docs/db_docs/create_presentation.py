import sys
import os
import pptx
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def build_presentation(output_path):
    prs = pptx.Presentation()
    # 16:9 Widescreen
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Color Palette (OSU Beaver Theme + Modern Tech)
    COLOR_BG_DARK = RGBColor(26, 26, 28)       # Deep charcoal/black
    COLOR_BG_LIGHT = RGBColor(248, 249, 250)   # Clean light gray
    COLOR_ORANGE = RGBColor(215, 63, 9)        # OSU Beaver Orange
    COLOR_WHITE = RGBColor(255, 255, 255)
    COLOR_TEXT_PRIMARY = RGBColor(33, 37, 41)
    COLOR_TEXT_MUTED = RGBColor(108, 117, 125)
    COLOR_CARD_BG = RGBColor(255, 255, 255)
    COLOR_CARD_BORDER = RGBColor(222, 226, 230)
    COLOR_GREEN = RGBColor(40, 167, 69)
    COLOR_RED = RGBColor(220, 53, 69)
    COLOR_ACCENT_BG = RGBColor(254, 243, 239)  # Light orange tint

    def add_header(slide, title_text, category="OPENBEAVS V2 ARCHITECTURE"):
        # Category / breadcrumb
        cat_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.7), Inches(0.35))
        tf_cat = cat_box.text_frame
        tf_cat.word_wrap = True
        tf_cat.margin_left = tf_cat.margin_top = tf_cat.margin_right = tf_cat.margin_bottom = 0
        p_cat = tf_cat.paragraphs[0]
        p_cat.text = category.upper()
        p_cat.font.size = Pt(10)
        p_cat.font.bold = True
        p_cat.font.color.rgb = COLOR_ORANGE

        # Main Slide Title
        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.7), Inches(11.7), Inches(0.6))
        tf_title = title_box.text_frame
        tf_title.word_wrap = True
        tf_title.margin_left = tf_title.margin_top = tf_title.margin_right = tf_title.margin_bottom = 0
        p_title = tf_title.paragraphs[0]
        p_title.text = title_text
        p_title.font.size = Pt(23)
        p_title.font.bold = True
        p_title.font.color.rgb = COLOR_TEXT_PRIMARY

    # =========================================================================
    # SLIDE 1: Title Slide (Dark Theme)
    # =========================================================================
    slide1 = prs.slides.add_slide(blank_layout)
    bg1 = slide1.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
    bg1.fill.solid()
    bg1.fill.fore_color.rgb = COLOR_BG_DARK
    bg1.line.fill.background()

    # Decorative Orange Stripe
    stripe = slide1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.8), Inches(0.12), Inches(3.2))
    stripe.fill.solid()
    stripe.fill.fore_color.rgb = COLOR_ORANGE
    stripe.line.fill.background()

    # Title & Subtitle Box
    t_box1 = slide1.shapes.add_textbox(Inches(1.2), Inches(1.7), Inches(11.0), Inches(3.4))
    tf1 = t_box1.text_frame
    tf1.word_wrap = True

    p0 = tf1.paragraphs[0]
    p0.text = "OREGON STATE UNIVERSITY CS CAPSTONE • PROJECT 39"
    p0.font.size = Pt(12)
    p0.font.bold = True
    p0.font.color.rgb = COLOR_ORANGE
    p0.space_after = Pt(14)

    p1 = tf1.add_paragraph()
    p1.text = "OpenBeavs v2: Database Simplification"
    p1.font.size = Pt(36)
    p1.font.bold = True
    p1.font.color.rgb = COLOR_WHITE
    p1.space_after = Pt(12)

    p2 = tf1.add_paragraph()
    p2.text = "Transitioning from 20+ Legacy Tables to a Lean 5-Table A2A Platform"
    p2.font.size = Pt(18)
    p2.font.color.rgb = RGBColor(180, 185, 190)

    # Footer Card on Slide 1
    f_box = slide1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.2), Inches(5.6), Inches(10.9), Inches(1.0))
    f_box.fill.solid()
    f_box.fill.fore_color.rgb = RGBColor(38, 38, 42)
    f_box.line.color.rgb = RGBColor(60, 60, 65)
    tf_f = f_box.text_frame
    tf_f.margin_left = Inches(0.3)
    p_f = tf_f.paragraphs[0]
    p_f.text = "Team: Jiya Pradhan • Faith Nambasa • Dean R Leon • Pedro P Blanco Jaramillo"
    p_f.font.size = Pt(12)
    p_f.font.color.rgb = RGBColor(220, 225, 230)
    p_f2 = tf_f.add_paragraph()
    p_f2.text = "Sponsor: John V. Sweet (john.sweet@oregonstate.edu)"
    p_f2.font.size = Pt(11)
    p_f2.font.color.rgb = COLOR_ORANGE

    # =========================================================================
    # SLIDE 2: The Problem: Legacy V1 Bloat
    # =========================================================================
    slide2 = prs.slides.add_slide(blank_layout)
    bg2 = slide2.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
    bg2.fill.solid()
    bg2.fill.fore_color.rgb = COLOR_BG_LIGHT
    bg2.line.fill.background()

    add_header(slide2, "The Problem: What We Inherited from Open WebUI", "CURRENT STATE & PAIN POINTS")

    cards_data_s2 = [
        ("20+ Inherited Tables (80% Unused)",
         "Open WebUI was built as a local ChatGPT clone with Slack-like channels, Python code valves, memory snippets, and prompt tools that have zero relevance to our A2A agent hub.",
         COLOR_RED),
        ("Dual ORMs & Fragile Migrations",
         "The backend ran both Peewee AND SQLAlchemy simultaneously. Startup migration wrappers and regex connection string hacks caused cold-start delays and connection leaks.",
         COLOR_RED),
        ("Monolithic JSON Chat Blobs",
         "Entire conversations were saved as a single JSON object (or encrypted blob). Appending one turn required loading, parsing, and re-encrypting the entire conversation history.",
         COLOR_RED),
        ("Conflicting Auth & Local Vector Stores",
         "Legacy had local email/password auth (triggering passlib bcrypt bugs) and ChromaDB/PgVector in the hub, when downstream A2A agents already handle their own embeddings.",
         COLOR_RED)
    ]

    card_w = Inches(5.6)
    card_h = Inches(2.3)
    positions_s2 = [
        (Inches(0.8), Inches(1.5)),
        (Inches(6.9), Inches(1.5)),
        (Inches(0.8), Inches(4.1)),
        (Inches(6.9), Inches(4.1))
    ]

    for i, (title, desc, tag_color) in enumerate(cards_data_s2):
        x, y = positions_s2[i]
        card = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, card_w, card_h)
        card.fill.solid()
        card.fill.fore_color.rgb = COLOR_CARD_BG
        card.line.color.rgb = COLOR_CARD_BORDER

        # Colored indicator strip
        ind = slide2.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y + Inches(0.15), Inches(0.08), Inches(0.45))
        ind.fill.solid()
        ind.fill.fore_color.rgb = tag_color
        ind.line.fill.background()

        tb = slide2.shapes.add_textbox(x + Inches(0.25), y + Inches(0.15), card_w - Inches(0.5), card_h - Inches(0.3))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_top = tf.margin_left = tf.margin_right = tf.margin_bottom = 0

        p_t = tf.paragraphs[0]
        p_t.text = title
        p_t.font.size = Pt(15)
        p_t.font.bold = True
        p_t.font.color.rgb = COLOR_TEXT_PRIMARY
        p_t.space_after = Pt(8)

        p_d = tf.add_paragraph()
        p_d.text = desc
        p_d.font.size = Pt(12)
        p_d.font.color.rgb = COLOR_TEXT_MUTED

    # Bottom takeaway pill
    pill2 = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(6.6), Inches(11.7), Inches(0.5))
    pill2.fill.solid()
    pill2.fill.fore_color.rgb = COLOR_ACCENT_BG
    pill2.line.color.rgb = COLOR_ORANGE
    tf_p2 = pill2.text_frame
    tf_p2.margin_left = Inches(0.25)
    p_p2 = tf_p2.paragraphs[0]
    p_p2.text = "Key Takeaway: v1's database was designed for a different product. Carrying it forward creates massive technical debt."
    p_p2.font.size = Pt(11)
    p_p2.font.bold = True
    p_p2.font.color.rgb = COLOR_ORANGE

    # =========================================================================
    # SLIDE 3: The Proposed Solution: 5 Clean Tables + Visual ERD
    # =========================================================================
    slide3 = prs.slides.add_slide(blank_layout)
    bg3 = slide3.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
    bg3.fill.solid()
    bg3.fill.fore_color.rgb = COLOR_BG_LIGHT
    bg3.line.fill.background()

    add_header(slide3, "The Solution: A Streamlined 5-Table Schema", "TARGET ARCHITECTURE FOR OPENBEAVS V2")

    # Left Column: The 5 Table Definitions (Width = 6.4 inches)
    tables_data = [
        ("agents", "Unified Agent Registry",
         "Replaces competing 'agent'/'registry_agent'. Stores A2A cards, endpoints, skills, and trust tiers (public, auth, privileged)."),
        ("users", "OSU SSO Identity",
         "Pure OIDC identity (ONID, email, role claim, KMS wrapping key). Zero local passwords or shadow accounts."),
        ("chats", "Conversation Sessions",
         "Tracks session turns, source domain (FAB origin), and 24-hour expiration TTL for guest conversations."),
        ("messages", "Relational Message Turns",
         "Sequential turns (user/agent/system). Enables fast streaming, indexing, and AES-256 envelope encryption."),
        ("tickets", "Feedback & Bug Reports",
         "Lightweight support tracking for user feedback, issues, and browser telemetry.")
    ]

    card_y_start = Inches(1.5)
    card_h_left = Inches(0.92)
    card_gap = Inches(0.08)

    for i, (tname, ttitle, tdesc) in enumerate(tables_data):
        cy = card_y_start + i * (card_h_left + card_gap)
        c_left = slide3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), cy, Inches(6.4), card_h_left)
        c_left.fill.solid()
        c_left.fill.fore_color.rgb = COLOR_CARD_BG
        c_left.line.color.rgb = COLOR_CARD_BORDER

        # Mini Orange Indicator
        c_ind = slide3.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), cy + Inches(0.12), Inches(0.06), card_h_left - Inches(0.24))
        c_ind.fill.solid()
        c_ind.fill.fore_color.rgb = COLOR_ORANGE
        c_ind.line.fill.background()

        tb_l = slide3.shapes.add_textbox(Inches(1.0), cy + Inches(0.08), Inches(6.1), card_h_left - Inches(0.16))
        tf_l = tb_l.text_frame
        tf_l.word_wrap = True
        tf_l.margin_top = tf_l.margin_left = tf_l.margin_right = tf_l.margin_bottom = 0

        p_t1 = tf_l.paragraphs[0]
        p_t1.text = f"{tname}  •  {ttitle}"
        p_t1.font.size = Pt(13)
        p_t1.font.bold = True
        p_t1.font.color.rgb = COLOR_TEXT_PRIMARY

        p_t2 = tf_l.add_paragraph()
        p_t2.text = tdesc
        p_t2.font.size = Pt(10)
        p_t2.font.color.rgb = COLOR_TEXT_MUTED

    # Right Column: Visual ERD Diagram
    erd_img_path = os.path.join(os.path.dirname(output_path), "database-erd.png")
    erd_bg = slide3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(7.4), Inches(1.5), Inches(5.1), Inches(4.9))
    erd_bg.fill.solid()
    erd_bg.fill.fore_color.rgb = RGBColor(18, 19, 21)  # Dark card to match ERD
    erd_bg.line.color.rgb = COLOR_CARD_BORDER

    if os.path.exists(erd_img_path):
        # Insert ERD image centered in right card
        slide3.shapes.add_picture(erd_img_path, Inches(7.55), Inches(1.6), width=Inches(4.8), height=Inches(4.7))
    else:
        # Fallback text box if image missing
        tb_erd = slide3.shapes.add_textbox(Inches(7.5), Inches(3.0), Inches(4.9), Inches(2.0))
        tb_erd.text_frame.word_wrap = True
        p_e = tb_erd.text_frame.paragraphs[0]
        p_e.text = "Visual ERD: users ──< chats ──< messages\nagents ──< chats\nusers ──< tickets"
        p_e.font.size = Pt(14)
        p_e.font.color.rgb = COLOR_WHITE

    # Bottom summary badge
    sum_box = slide3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(6.6), Inches(11.7), Inches(0.5))
    sum_box.fill.solid()
    sum_box.fill.fore_color.rgb = RGBColor(240, 243, 246)
    sum_box.line.color.rgb = COLOR_CARD_BORDER
    tf_sum = sum_box.text_frame
    tf_sum.margin_left = Inches(0.25)
    p_sum = tf_sum.paragraphs[0]
    p_sum.text = "Stack: Pure SQLAlchemy 2.0 (async/sync) + Alembic • PostgreSQL (Cloud SQL) / SQLite (Local Dev)"
    p_sum.font.size = Pt(10.5)
    p_sum.font.bold = True
    p_sum.font.color.rgb = COLOR_TEXT_PRIMARY

    # =========================================================================
    # SLIDE 4: Architecture & Operational Impact (V1 vs V2)
    # =========================================================================
    slide4 = prs.slides.add_slide(blank_layout)
    bg4 = slide4.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
    bg4.fill.solid()
    bg4.fill.fore_color.rgb = COLOR_BG_LIGHT
    bg4.line.fill.background()

    add_header(slide4, "Architectural & Operational Impact (v1 vs. v2)", "BENEFITS & METRIC COMPARISON")

    # Table comparison
    table_shape = slide4.shapes.add_table(6, 3, Inches(0.8), Inches(1.5), Inches(11.73), Inches(4.8))
    table = table_shape.table
    table.columns[0].width = Inches(2.73)
    table.columns[1].width = Inches(4.5)
    table.columns[2].width = Inches(4.5)

    headers = ["Dimension", "Legacy OpenBeavs (v1)", "OpenBeavs v2 (Target)"]
    for j, h in enumerate(headers):
        cell = table.cell(0, j)
        cell.fill.solid()
        cell.fill.fore_color.rgb = COLOR_BG_DARK
        p = cell.text_frame.paragraphs[0]
        p.text = h
        p.font.bold = True
        p.font.size = Pt(13)
        p.font.color.rgb = COLOR_WHITE

    rows_data = [
        ("Table Count", "20+ tables (80% dead weight)", "5 core tables (-75% schema footprint)"),
        ("ORM & Migrations", "Dual Peewee + SQLAlchemy with boot hacks", "Standard SQLAlchemy 2.0 + Alembic only"),
        ("Message Storage", "Monolithic JSON rewrite on every turn", "Append-only relational rows (streamable)"),
        ("Authentication", "Local passwords + passlib bcrypt bugs", "Native OSU Microsoft SSO + 24h Guest TTL"),
        ("Hub Vector DB", "ChromaDB + PgVector overhead in hub", "Zero (downstream A2A agents manage RAG)")
    ]

    for i, (dim, v1, v2) in enumerate(rows_data):
        row_idx = i + 1
        c0 = table.cell(row_idx, 0)
        c1 = table.cell(row_idx, 1)
        c2 = table.cell(row_idx, 2)

        bg_color = COLOR_CARD_BG if i % 2 == 0 else RGBColor(245, 247, 250)
        for c in (c0, c1, c2):
            c.fill.solid()
            c.fill.fore_color.rgb = bg_color

        p0 = c0.text_frame.paragraphs[0]
        p0.text = dim
        p0.font.bold = True
        p0.font.size = Pt(12)
        p0.font.color.rgb = COLOR_TEXT_PRIMARY

        p1 = c1.text_frame.paragraphs[0]
        p1.text = v1
        p1.font.size = Pt(12)
        p1.font.color.rgb = RGBColor(160, 40, 40)

        p2 = c2.text_frame.paragraphs[0]
        p2.text = v2
        p2.font.bold = True
        p2.font.size = Pt(12)
        p2.font.color.rgb = COLOR_GREEN

    # Bottom impact pill
    pill = slide4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(6.55), Inches(11.73), Inches(0.55))
    pill.fill.solid()
    pill.fill.fore_color.rgb = COLOR_ACCENT_BG
    pill.line.color.rgb = COLOR_ORANGE
    tf_pill = pill.text_frame
    tf_pill.margin_left = Inches(0.3)
    p_pill = tf_pill.paragraphs[0]
    p_pill.text = "Result: Eliminates Cloud Run boot latency (<1s cold starts) and removes thousands of lines of tech debt."
    p_pill.font.size = Pt(11)
    p_pill.font.bold = True
    p_pill.font.color.rgb = COLOR_ORANGE

    # =========================================================================
    # SLIDE 5: Migration Roadmap & Next Steps
    # =========================================================================
    slide5 = prs.slides.add_slide(blank_layout)
    bg5 = slide5.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
    bg5.fill.solid()
    bg5.fill.fore_color.rgb = COLOR_BG_LIGHT
    bg5.line.fill.background()

    add_header(slide5, "Implementation Roadmap & Next Steps", "EXECUTION PLAN FOR OPENBEAVS V2")

    steps_data = [
        ("1. Setup DB Core", "backend/app/db/",
         "Initialize modern SQLAlchemy 2.0 engine, async session factory, and Alembic migration directory."),
        ("2. Define Models", "backend/app/models/",
         "Implement the 5 clean declarative classes: Agent, User, Chat, Message, Ticket with proper indices."),
        ("3. Connect Chris", "backend/app/chris/",
         "Swap out hardcoded MOCK_AGENTS for live queries against the agents table; pass candidate agents to LLM."),
        ("4. Persist Chats", "POST /chris/message",
         "Store user turns and agent responses into chats and messages tables with session tracking and guest TTL.")
    ]

    sw_w = Inches(2.78)
    sw_h = Inches(4.3)
    sw_gap = Inches(0.2)

    for i, (stitle, spath, sdesc) in enumerate(steps_data):
        sx = Inches(0.8) + i * (sw_w + sw_gap)
        sy = Inches(1.5)

        scard = slide5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, sx, sy, sw_w, sw_h)
        scard.fill.solid()
        scard.fill.fore_color.rgb = COLOR_CARD_BG
        scard.line.color.rgb = COLOR_CARD_BORDER

        # Phase header
        sbar = slide5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, sx, sy, sw_w, Inches(1.1))
        sbar.fill.solid()
        sbar.fill.fore_color.rgb = COLOR_BG_DARK
        sbar.line.fill.background()

        tb_sh = slide5.shapes.add_textbox(sx + Inches(0.15), sy + Inches(0.12), sw_w - Inches(0.3), Inches(0.9))
        tf_sh = tb_sh.text_frame
        tf_sh.word_wrap = True
        tf_sh.margin_top = tf_sh.margin_left = tf_sh.margin_right = tf_sh.margin_bottom = 0

        p_st = tf_sh.paragraphs[0]
        p_st.text = stitle
        p_st.font.size = Pt(14)
        p_st.font.bold = True
        p_st.font.color.rgb = COLOR_WHITE

        p_sp = tf_sh.add_paragraph()
        p_sp.text = spath
        p_sp.font.size = Pt(11)
        p_sp.font.color.rgb = COLOR_ORANGE

        # Body
        tb_sb = slide5.shapes.add_textbox(sx + Inches(0.18), sy + Inches(1.3), sw_w - Inches(0.36), sw_h - Inches(1.5))
        tf_sb = tb_sb.text_frame
        tf_sb.word_wrap = True
        tf_sb.margin_top = tf_sb.margin_left = tf_sb.margin_right = tf_sb.margin_bottom = 0

        p_sd = tf_sb.paragraphs[0]
        p_sd.text = sdesc
        p_sd.font.size = Pt(12)
        p_sd.font.color.rgb = COLOR_TEXT_MUTED

    # Bottom Call to Action / Summary
    callout = slide5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(6.05), Inches(11.73), Inches(0.85))
    callout.fill.solid()
    callout.fill.fore_color.rgb = COLOR_BG_DARK
    callout.line.fill.background()

    tf_call = callout.text_frame
    tf_call.margin_left = Inches(0.3)
    p_c1 = tf_call.paragraphs[0]
    p_c1.text = "Summary: A clean database architecture eliminates technical debt early, giving the team a robust,"
    p_c1.font.size = Pt(11)
    p_c1.font.color.rgb = COLOR_WHITE
    p_c2 = tf_call.add_paragraph()
    p_c2.text = "scalable platform to build features rapidly across Winter and Spring terms."
    p_c2.font.size = Pt(11)
    p_c2.font.bold = True
    p_c2.font.color.rgb = COLOR_ORANGE

    prs.save(output_path)
    print(f"Presentation saved successfully to {output_path}")

if __name__ == "__main__":
    out_file = sys.argv[1] if len(sys.argv) > 1 else "OpenBeavs_v2_Database_Simplification.pptx"
    build_presentation(out_file)
